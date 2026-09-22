import json,re,datetime
from pathlib import Path
import openpyxl
root=Path(__file__).parent
w=openpyxl.load_workbook(root/'ColinsPhillipsBowl-2026-2027.xlsx',data_only=True)
players=['Bryan','Ed','Mike','Kevin']
weeks=[]
for n in range(1,19):
 s=next(s for s in w if s.title.strip()==f'Week {n}')
 games=[]
 for row in range(2,22):
  m=s.cell(row,1).value
  if not m: continue
  teams=[t.replace('Louisvile','Louisville').strip() for t in re.split(r'\s+@\s+|\s+at\s+',re.sub(r'\s*\([^)]*\)','',m).strip())]
  games.append(dict(id=f'{n}-{row}',matchup=m.strip(),teams=teams,picks={p:s.cell(row,i+2).value or '' for i,p in enumerate(players)},winner=s.cell(row,6).value or '',score=s.cell(row,7).value or '',kickoff=''))
 e=w['Earnings']
 weeks.append(dict(number=n,games=games,totalPoints={p:s.cell(22,i+2).value for i,p in enumerate(players)},recordedTotals={p:s.cell(23,i+2).value if games else None for i,p in enumerate(players)},winner=e.cell(n+1,3).value or '',earnings=e.cell(n+1,4).value or 0,note=e.cell(n+1,5).value or '',dues={p:e.cell(n+1,c).value or 0 for p,c in [('Ed',8),('Kevin',9),('Bryan',10),('Mike',11)]},teamEarnings={'Young Guns':e.cell(n+24,3).value or 0,'Old Farts':e.cell(n+24,4).value or 0},actualTotal=34 if n==2 else None,locked=n<3,deadlineLocal=(datetime.datetime(2026,9,23,23,59)+datetime.timedelta(weeks=n-3)).isoformat(timespec="minutes")))
data=dict(season='2026–2027',source='ColinsPhillipsBowl-2026-2027.xlsx',players=players,teams={'Young Guns':['Kevin','Mike'],'Old Farts':['Ed','Bryan']},weeks=weeks,finalSchedules={str(n):[w['Final Schedules'].cell(r,n-13).value for r in range(2,22) if w['Final Schedules'].cell(r,n-13).value] for n in range(14,19)})
(root/'pool-site/lib/seed.json').write_text(json.dumps(data,indent=2))
print(f'Imported {len(weeks)} weeks, {sum(len(x["games"]) for x in weeks)} games.')
